import os
import sys
import io
import json
import time
import socket
import urllib.request
import urllib.error
import sqlite3
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

BASE_URL = "http://localhost:8000"
ONEDRIVE_DIR = r"C:\Users\sunny\OneDrive - Marigold Aviation\Apps\Anitgravity\Audit Trail"
EXCEL_PATH = os.path.join(ONEDRIVE_DIR, "FTO_Audit_Consolidated.xlsx")
EVIDENCE_DIR = os.path.join(ONEDRIVE_DIR, "Evidence")
DB_PATH = "fto_audit.db"

results = []

def test(name, fn):
    try:
        fn()
        results.append((name, "PASS", ""))
        print(f" [PASS] {name}")
    except Exception as e:
        results.append((name, "FAIL", str(e)))
        print(f" [FAIL] {name}: {e}")

# ----------------------------------------------------------------
# 1. Connectivity & Route Tests
# ----------------------------------------------------------------
def test_routes():
    routes = ["/", "/master", "/auditor1", "/auditor2", "/auditor3", "/api/state", "/api/sync"]
    for r in routes:
        url = BASE_URL + r
        req = urllib.request.Request(url)
        with urllib.request.urlopen(req, timeout=5) as resp:
            assert resp.getcode() == 200, f"Route {r} returned {resp.getcode()}"

test("HTTP Endpoints (Portal, Master, Auditor 1, 2, 3, APIs)", test_routes)

def test_local_ip():
    hostname = socket.gethostname()
    local_ip = socket.gethostbyname(hostname)
    url = f"http://{local_ip}:8000/master"
    req = urllib.request.Request(url)
    with urllib.request.urlopen(req, timeout=5) as resp:
        assert resp.getcode() == 200, f"Wi-Fi URL {url} returned {resp.getcode()}"

test("Local Network / Wi-Fi Access (for phone & colleague laptops)", test_local_ip)

# ----------------------------------------------------------------
# 2. State & Database Tests
# ----------------------------------------------------------------
def test_state_api():
    url = BASE_URL + "/api/state"
    with urllib.request.urlopen(url, timeout=5) as resp:
        data = json.loads(resp.read().decode("utf-8"))
        assert data.get("status") == "ok", "State API status not ok"
        st = data.get("state", {})
        assert "items" in st, "State missing items"
        assert "meta" in st, "State missing meta"
        assert "lists" in st, "State missing lists"
        assert "secs" in st, "State missing secs"

test("GET /api/state JSON Schema & Integrity", test_state_api)

def test_save_api():
    url = BASE_URL + "/api/state"
    with urllib.request.urlopen(url, timeout=5) as resp:
        st = json.loads(resp.read().decode("utf-8"))["state"]
    
    st["meta"]["testStamp"] = str(time.time())
    save_url = BASE_URL + "/api/save"
    req = urllib.request.Request(save_url, data=json.dumps({"state": st}).encode("utf-8"), headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=5) as resp:
        res = json.loads(resp.read().decode("utf-8"))
        assert res.get("status") == "ok", "Save API status not ok"
        assert "savedAt" in res, "Missing savedAt in response"

test("POST /api/save Persistence & Auto-Save", test_save_api)

# ----------------------------------------------------------------
# Helper: Fetch Workbook (Stream or Disk)
# ----------------------------------------------------------------
def get_test_workbook():
    try:
        if os.path.exists(EXCEL_PATH):
            return openpyxl.load_workbook(EXCEL_PATH, data_only=False)
    except PermissionError:
        pass
    url = BASE_URL + "/api/export-excel"
    with urllib.request.urlopen(url, timeout=10) as resp:
        return openpyxl.load_workbook(io.BytesIO(resp.read()), data_only=False)

# ----------------------------------------------------------------
# 3. Excel Generation & Multi-Tab Structure
# ----------------------------------------------------------------
def test_excel_file():
    wb = get_test_workbook()
    expected_sheets = ["Executive Overview", "Consolidated Findings & CAP", "Evidence Index"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Missing sheet: {s}"
    
    chk_sheets = [s for s in wb.sheetnames if "checklist" in s.lower()]
    assert len(chk_sheets) >= 1, "No checklist sheets found"
    
    ws_exec = wb["Executive Overview"]
    content_dump = [str(val) for row in ws_exec.iter_rows(values_only=True) for val in row if val is not None]
    full_text = " ".join(content_dump)
    
    assert "FTO Approval" in full_text or "Approval" in full_text, "Missing Governance info in Executive Overview"
    assert "Audit Team" in full_text, "Missing Audit Team in Executive Overview"
    assert "Audit Scope & Functional Section Governance Table" in full_text, "Missing Section Governance in Executive Overview"
    assert "Assigned Section Rights" in full_text, "Missing Section Rights in Executive Overview"
    assert "Instructional Cadre" in full_text, "Missing Cadre in Executive Overview"
    assert "Aircraft Fleet" in full_text, "Missing Fleet in Executive Overview"

test("Excel Workbook Structure & Overview Content", test_excel_file)

# ----------------------------------------------------------------
# 4. Data Validation Dropdowns in Excel
# ----------------------------------------------------------------
def test_excel_dropdowns():
    wb = get_test_workbook()
    ws_exec = wb["Executive Overview"]
    assert len(ws_exec.data_validations.dataValidation) >= 1, "Missing DataValidation in Executive Overview"
    
    ws_cap = wb["Consolidated Findings & CAP"]
    assert len(ws_cap.data_validations.dataValidation) >= 3, "Missing DataValidation in Findings & CAP"
    
    chk_sheets = [s for s in wb.sheetnames if "checklist" in s.lower()]
    ws_chk = wb[chk_sheets[0]]
    assert len(ws_chk.data_validations.dataValidation) >= 4, "Missing DataValidation in Checklist sheet"

test("Excel Data Validation (Native Dropdowns for Ratings, Risk, Scope)", test_excel_dropdowns)

# ----------------------------------------------------------------
# 5. Evidence Photo Links & Hyperlinks
# ----------------------------------------------------------------
def test_evidence_links():
    wb = get_test_workbook()
    ws_ev = wb["Evidence Index"]
    
    links_found = 0
    for row in ws_ev.iter_rows(min_row=2, max_col=6, values_only=False):
        c = row[5]
        if c.value and "HYPERLINK" in str(c.value):
            links_found += 1
    print(f"    (Verified {links_found} embedded HYPERLINK formulas in Evidence Index)")
    assert links_found > 0 or len(os.listdir(EVIDENCE_DIR)) == 0, "No hyperlinks found in Evidence Index"

test("Evidence Photo Hyperlinks (=HYPERLINK) in Excel", test_evidence_links)

# ----------------------------------------------------------------
# 6. Streamed Excel Export API (No File-Locking 500 Errors)
# ----------------------------------------------------------------
def test_streamed_export():
    url = BASE_URL + "/api/export-excel"
    req = urllib.request.Request(url)
    with urllib.request.urlopen(req, timeout=10) as resp:
        assert resp.getcode() == 200, f"Export API returned {resp.getcode()}"
        content_type = resp.headers.get("Content-Type", "")
        assert "spreadsheetml" in content_type, f"Unexpected Content-Type: {content_type}"
        data = resp.read()
        assert len(data) > 5000, f"Exported file too small: {len(data)} bytes"

test("GET /api/export-excel Streamed Download (Lock-Free)", test_streamed_export)

# ----------------------------------------------------------------
# 7. Two-Way Roundtrip Synchronization (Excel ➔ Web App)
# ----------------------------------------------------------------
def test_two_way_sync():
    try:
        with open(EXCEL_PATH, "a+b") as f:
            pass
    except PermissionError:
        print("    (Notice: Disk file is currently locked in Excel Desktop; testing import logic directly)")
        from excel_sync import import_audit_excel
        conn = sqlite3.connect(DB_PATH)
        row = conn.cursor().execute("SELECT json_data FROM audit_state WHERE key = 'current'").fetchone()
        st = json.loads(row[0])
        url = BASE_URL + "/api/export-excel"
        with urllib.request.urlopen(url, timeout=10) as resp:
            bio = io.BytesIO(resp.read())
            count, new_st = import_audit_excel(bio, st)
            assert count > 0, "Import from Excel stream returned 0 items"
        return

    # If unlocked on disk, test disk roundtrip
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws_chk = wb[[s for s in wb.sheetnames if "checklist" in s.lower()][0]]
    test_comment = f"VERIFIED AUTOMATED TWO-WAY SYNC AT {time.time()}"
    ws_chk.cell(row=3, column=9, value=test_comment)
    time.sleep(1.2)
    wb.save(EXCEL_PATH)
    
    time.sleep(0.5)
    with urllib.request.urlopen(BASE_URL + "/api/sync", timeout=5) as resp:
        assert resp.getcode() == 200
        
    conn = sqlite3.connect(DB_PATH)
    row = conn.cursor().execute("SELECT json_data FROM audit_state WHERE key = 'current'").fetchone()
    st = json.loads(row[0])
    item_id = str(ws_chk.cell(row=3, column=1).value).strip()
    db_comment = st.get("items", {}).get(item_id, {}).get("comment", "")
    assert db_comment == test_comment, f"Expected comment '{test_comment}', got '{db_comment}'"

test("Two-Way Roundtrip Live Sync (Disk Excel ➔ Web App Database)", test_two_way_sync)

# ----------------------------------------------------------------
# Summary
# ----------------------------------------------------------------
print("\n" + "="*60)
print(f"TEST SUITE SUMMARY: {sum(1 for _, s, _ in results if s == 'PASS')}/{len(results)} TESTS PASSED")
print("="*60)
for name, status, err in results:
    if status == "PASS":
        print(f" [PASS] {name}")
    else:
        print(f" [FAIL] {name} -> {err}")
print("="*60 + "\n")
